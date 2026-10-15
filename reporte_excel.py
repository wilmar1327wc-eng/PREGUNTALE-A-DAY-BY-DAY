# -*- coding: utf-8 -*-
"""
reporte_excel.py — Mismo informe ejecutivo que el PDF, pero en Excel (.xlsx).

A diferencia del PDF, aquí van los NÚMEROS de verdad (no texto), para que el
usuario pueda filtrar, pivotear y graficar. Identidad job&talent (encabezado
azul profundo + acentos teal). Sin dependencias de la app: recibe todo calculado.

Pestañas:
  - Resumen    : KPI + variación vs período anterior + hallazgos.
  - Desglose   : tabla filtrable por dimensión (con barras de datos y total).
  - Tendencia  : serie en el tiempo + gráfico de línea nativo de Excel.
"""

import io
import datetime as _dt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter

# ---- paleta (clara, para Excel) ----
NAVY = "0E2A52"
TEAL = "07B6A6"
TEAL_D = "059B8E"
INK = "15233B"
MUTED = "5C6E88"
TINT = "E9F7F5"
LINEC = "D9E1EC"
BAD = "C0453A"

FONT = "Arial"
_HAIR = Side(style="thin", color=LINEC)
BORDER = Border(left=_HAIR, right=_HAIR, top=_HAIR, bottom=_HAIR)


def _numfmt(fmt):
    if fmt == "pct":
        return "0.0%"
    if fmt == "money":
        return '"$"\\ #,##0'
    return "#,##0"


def _band(ws, text, sub, fecha_txt, ncols):
    """Encabezado de marca (banda azul) sobre las primeras filas."""
    last = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last}1")
    ws.merge_cells(f"A2:{last}2")
    ws.merge_cells(f"A3:{last}3")
    c = ws["A1"]
    c.value = "Centro de Inteligencia de Negocios  ·  job&talent"
    c.font = Font(name=FONT, size=12, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(vertical="center", indent=1)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    for col in range(1, ncols + 1):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor=NAVY)
        ws.cell(row=2, column=col).fill = PatternFill("solid", fgColor=TEAL)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 4
    t = ws["A3"]
    t.value = text
    t.font = Font(name=FONT, size=15, bold=True, color=INK)
    t.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[3].height = 24
    ws.merge_cells(f"A4:{last}4")
    s = ws["A4"]
    s.value = f"{sub}    ·    Informe generado: {fecha_txt}"
    s.font = Font(name=FONT, size=10, color=MUTED)
    s.alignment = Alignment(vertical="center", indent=1)


def construir_xlsx(*, indicador, cliente, periodo, metrica_label, fmt, tipo,
                   valor, base, var=None, var_cls="flat",
                   filas, dim_label, serie=(), insights=(), fecha_txt=""):
    """
    filas: lista de (dimension:str, valor:float, base:int)
    serie: lista de (fecha:date, valor:float)
    insights: lista de str
    """
    nf = _numfmt(fmt)
    wb = Workbook()

    # ============ Pestaña 1: Resumen ============
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 24
    _band(ws, f"Informe de {indicador}", f"{cliente}  ·  {periodo}", fecha_txt, 3)

    r = 6
    ws.cell(r, 1, metrica_label).font = Font(name=FONT, size=10, color=MUTED)
    kc = ws.cell(r + 1, 1, float(valor))
    kc.number_format = nf
    kc.font = Font(name=FONT, size=26, bold=True, color=NAVY)
    kc.alignment = Alignment(horizontal="left")
    ws.cell(r + 2, 1, f"sobre {int(base):,} registros".replace(",", ".")).font = Font(name=FONT, size=9, color=MUTED)
    if var is not None:
        vtxt = f"{var * 100:+.1f}%".replace(".", ",") + "  vs. período anterior"
        vcolor = {"good": TEAL_D, "bad": BAD}.get(var_cls, MUTED)
        vc = ws.cell(r + 3, 1, vtxt)
        vc.font = Font(name=FONT, size=11, bold=True, color=vcolor)
        vc.alignment = Alignment(horizontal="left")
    # marco teal del bloque KPI
    for rr in range(r, r + 4):
        ws.cell(rr, 1).border = Border(left=Side(style="thick", color=TEAL))

    # Hallazgos
    hr = r + 6
    ws.cell(hr, 1, "Hallazgos").font = Font(name=FONT, size=12, bold=True, color=INK)
    for i, t in enumerate(insights, 1):
        cell = ws.cell(hr + i, 1, f"•  {t}")
        cell.font = Font(name=FONT, size=10, color=INK)
        ws.merge_cells(start_row=hr + i, start_column=1, end_row=hr + i, end_column=3)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[hr + i].height = 30

    # ============ Pestaña 2: Desglose ============
    wd = wb.create_sheet("Desglose")
    wd.sheet_view.showGridLines = False
    wd.column_dimensions["A"].width = 34
    wd.column_dimensions["B"].width = 20
    wd.column_dimensions["C"].width = 16
    _band(wd, f"Desglose por {dim_label}", f"{indicador}  ·  {cliente}  ·  {periodo}", fecha_txt, 3)

    hrow = 6
    heads = [dim_label.capitalize(), metrica_label, "Registros"]
    for j, h in enumerate(heads, 1):
        c = wd.cell(hrow, j, h)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal=("left" if j == 1 else "right"), vertical="center", indent=1)
        c.border = BORDER
    wd.row_dimensions[hrow].height = 20

    first = hrow + 1
    for i, (dim, v, b) in enumerate(filas):
        rr = first + i
        a = wd.cell(rr, 1, str(dim)); a.font = Font(name=FONT, size=10, color=INK)
        a.alignment = Alignment(indent=1); a.border = BORDER
        cv = wd.cell(rr, 2, float(v)); cv.number_format = nf
        cv.font = Font(name=FONT, size=10, bold=True, color=INK); cv.border = BORDER
        cb = wd.cell(rr, 3, int(b)); cb.number_format = "#,##0"
        cb.font = Font(name=FONT, size=10, color=MUTED); cb.border = BORDER
        if i % 2 == 1:
            for j in range(1, 4):
                wd.cell(rr, j).fill = PatternFill("solid", fgColor="F5F8FB")
    last = first + len(filas) - 1

    # Total (suma real para sum/count; "General" ponderado del cubo para ratio/avg)
    if last >= first:
        tr = last + 1
        lbl = wd.cell(tr, 1, "TOTAL" if tipo in ("sum", "count") else "General (del cubo)")
        lbl.font = Font(name=FONT, size=10, bold=True, color=NAVY)
        lbl.alignment = Alignment(indent=1)
        tc = wd.cell(tr, 2)
        if tipo in ("sum", "count"):
            tc.value = f"=SUM(B{first}:B{last})"
        else:
            tc.value = float(valor)
        tc.number_format = nf
        tc.font = Font(name=FONT, size=10, bold=True, color=NAVY)
        tb = wd.cell(tr, 3, f"=SUM(C{first}:C{last})")
        tb.number_format = "#,##0"
        tb.font = Font(name=FONT, size=10, bold=True, color=NAVY)
        for j in range(1, 4):
            wd.cell(tr, j).border = Border(top=Side(style="thin", color=NAVY))

        # barras de datos (si la versión de openpyxl las soporta) + filtro + congelar
        try:
            rule = DataBarRule(start_type="min", end_type="max", color=TEAL, showValue=True)
            wd.conditional_formatting.add(f"B{first}:B{last}", rule)
        except Exception:
            pass
        wd.auto_filter.ref = f"A{hrow}:C{last}"
        wd.freeze_panes = f"A{first}"

    # ============ Pestaña 3: Tendencia ============
    def _cd(x):
        # Deja una fecha limpia (sin hora ni zona horaria); openpyxl no acepta tz.
        try:
            if hasattr(x, "to_pydatetime"):   # pandas.Timestamp
                x = x.to_pydatetime()
            if isinstance(x, _dt.datetime):
                return x.date()
        except Exception:
            pass
        return x
    serie = [(_cd(f), float(v)) for f, v in (serie or []) if v == v]
    if len(serie) >= 2:
        wt = wb.create_sheet("Tendencia")
        wt.sheet_view.showGridLines = False
        wt.column_dimensions["A"].width = 16
        wt.column_dimensions["B"].width = 20
        _band(wt, "Tendencia del período", f"{indicador}  ·  {cliente}  ·  {periodo}", fecha_txt, 6)
        wt.cell(6, 1, "Fecha").font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        wt.cell(6, 2, metrica_label).font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        for j in (1, 2):
            wt.cell(6, j).fill = PatternFill("solid", fgColor=NAVY)
            wt.cell(6, j).alignment = Alignment(horizontal=("left" if j == 1 else "right"), indent=1)
        f0 = 7
        for i, (f, v) in enumerate(serie):
            rr = f0 + i
            cf = wt.cell(rr, 1, f)
            cf.number_format = "dd/mm/yyyy"
            cf.font = Font(name=FONT, size=10, color=INK)
            cv = wt.cell(rr, 2, v); cv.number_format = nf
            cv.font = Font(name=FONT, size=10, color=INK)
        flast = f0 + len(serie) - 1

        try:
            ch = LineChart()
            ch.title = f"{metrica_label} · {cliente}"
            ch.style = 2
            ch.height = 7.5
            ch.width = 18
            data = Reference(wt, min_col=2, min_row=6, max_row=flast)
            cats = Reference(wt, min_col=1, min_row=f0, max_row=flast)
            ch.add_data(data, titles_from_data=True)
            ch.set_categories(cats)
            ch.legend = None
            s0 = ch.series[0]
            s0.smooth = True
            s0.graphicalProperties.line.solidFill = TEAL
            s0.graphicalProperties.line.width = 28000  # EMU (~2.2pt)
            wt.add_chart(ch, f"A{flast + 3}")
        except Exception:
            pass   # si el gráfico falla en alguna versión, igual queda la tabla de datos

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()
